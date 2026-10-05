import collections
import hashlib
import json
import re
import unicodedata
from pathlib import Path
import openpyxl

SOURCE = Path(r'C:\Users\Francisco Fernandez\Downloads\AUTOEVALUACION 2025 (respuestas) (1).xlsx')
OUT = Path(__file__).parent
sheet = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True).worksheets[0]
raw = list(sheet.values)
width = max(map(len, raw))
rows = [list(r) + [None] * (width-len(r)) for r in raw]
headers, data = rows[0], rows[1:]
def clean(s):
    s = re.sub(r'\s+', ' ', str(s or '')).strip()
    match = re.fullmatch(r'\[(.*)\](?:\s+\d+)?', s)
    return (match.group(1) if match else s).strip().rstrip('.').strip()
def key(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s.casefold()) if not unicodedata.combining(c))
names = {'ENFERMERIA':'Enfermería','LIMPIEZA':'Limpieza','MANTENIMIENTO':'Mantenimiento','QUIROFANO':'Quirófano','SISTEMAS':'Sistemas','FACTURACIÓN':'Facturación','ENFERMERIA - AUXILIAR DE TRANSPORTE':'Enfermería - Auxiliar de transporte','RECEPCIÓN / CAJERO':'Recepción / Cajero','AUXILIAR DE SERVICIOS':'Auxiliar de servicios'}
counts = collections.Counter(r[9] for r in data if r[9])
sectors = [dict(id=f'S{i:02}',nombre=names[n],nombre_original=n,respuestas=counts[n]) for i,n in enumerate(counts,1)]
sector_id = {s['nombre_original']:s['id'] for s in sectors}
competencies=[]; origins=[]; excluded=[]; by_key={}; links=set()
for i,h in enumerate(headers):
    col=openpyxl.utils.get_column_letter(i+1)
    populated=[r for r in data if r[i] is not None]
    text=clean(h)
    if i<4 or i==9 or col=='BI' or not text or re.fullmatch(r'Fila \d+',text):
        reason='identificación' if i<4 else 'sector' if i==9 else 'campo sin significado definido; conservar aparte' if col=='BI' else 'marcador vacío sin competencia'
        excluded.append(dict(columna=col,encabezado=h,respuestas=len(populated),motivo=reason))
        continue
    kind='generica' if 4<=i<=8 else 'especifica'
    k=(kind,key(text))
    if k not in by_key:
        title,sep,description=text.partition(':')
        if title=='Excelencia Operativa': title='Excelencia operativa'
        item=dict(id=f'{"G" if kind=="generica" else "E"}{sum(c["tipo"]==kind for c in competencies)+1:03}',tipo=kind,nombre=title.strip(),descripcion=description.strip() if sep else None)
        competencies.append(item); by_key[k]=item
    item=by_key[k]
    observed=collections.Counter(r[9] for r in populated if r[9])
    suggestion='Administración (inferido por contenido)' if 30<=i<=34 else 'RR. HH. / Liquidación de sueldos (inferido por contenido)' if not observed and kind=='especifica' else None
    origins.append(dict(competencia_id=item['id'],hoja=sheet.title,columna=col,encabezado_original=h,respuestas=len(populated),sectores_observados={sector_id[s]:n for s,n in observed.items()},area_sugerida=suggestion))
    if kind=='especifica':
        links.update((sector_id[s],item['id']) for s in observed)
templates=[]; items=[]
for sector in sectors:
    tid='P'+sector['id'][1:]
    templates.append(dict(id=tid,sector_id=sector['id'],puesto_id=None,nombre='Autoevaluación 2025 - '+sector['nombre'],version=1,estado='borrador',anio=2025))
    selected=[c for c in competencies if c['tipo']=='generica' or (sector['id'],c['id']) in links]
    for order,c in enumerate(selected,1):
        items.append(dict(plantilla_id=tid,competencia_id=c['id'],orden=order,peso=None,obligatorio=None))
missing=[]
for rownum,r in enumerate(data,2):
    if not r[9]: continue
    for o in origins:
        if o['competencia_id'].startswith('G') or sector_id[r[9]] in o['sectores_observados']:
            idx=openpyxl.utils.column_index_from_string(o['columna'])-1
            if r[idx] is None: missing.append(dict(fila=rownum,columna=o['columna'],sector_id=sector_id[r[9]],competencia_id=o['competencia_id']))
duplicates=[dict(competencia_id=c['id'],columnas=[o['columna'] for o in origins if o['competencia_id']==c['id']]) for c in competencies if sum(o['competencia_id']==c['id'] for o in origins)>1]
unassigned=[c['id'] for c in competencies if c['tipo']=='especifica' and not any(cid==c['id'] for _,cid in links)]
observed_scale=sorted({str(r[i]) for r in data for i in range(4,width) if (4<=i<=8 or any(o['columna']==openpyxl.utils.get_column_letter(i+1) for o in origins)) and r[i] is not None})
payload=dict(fuente=dict(archivo=SOURCE.name,hoja=sheet.title,sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest()),competencias_genericas=[c for c in competencies if c['tipo']=='generica'],competencias_especificas=[c for c in competencies if c['tipo']=='especifica'],sectores=sectors,puestos=[],sector_competencias=[dict(sector_id=s,competencia_id=c,evidencia='coincidencia de respuestas con columna J') for s,c in sorted(links)],plantillas_evaluacion=templates,plantilla_competencias=items,origen_columnas=origins,columnas_excluidas_del_catalogo=excluded,duplicados_unificados=duplicates,competencias_sin_sector_confirmado=unassigned,respuestas_faltantes=missing,escala_observada=observed_scale)
(OUT/'catalogo_normalizado.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
assert len(origins)+len(excluded)==width
assert len({c['id'] for c in competencies})==len(competencies)
assert all(len([x for x in items if x['plantilla_id']==t['id']])==10 for t in templates)
assert len(items)==len({(x['plantilla_id'],x['competencia_id']) for x in items})
assert all(x['competencia_id'] in {c['id'] for c in competencies} for x in items)
lines=['# Autoevaluación 2025: catálogo depurado','',f'Fuente: {SOURCE.name}, hoja «{sheet.title}».','',f'83 columnas; 133 respuestas con sector; 42 filas completamente vacías; 9 filas adicionales sin marca temporal ni sector (176–184), con datos solo en C. No se interpretan como evaluaciones.','',f'Catálogo: 5 genéricas y {len(payload["competencias_especificas"])} específicas únicas. De las específicas, 45 tienen sector observado y {len(unassigned)} quedan pendientes de asignación.','', '## Competencias genéricas (E–I)','']
for c in payload['competencias_genericas']: lines.append(f'- **{c["nombre"]}**: {c["descripcion"]}')
lines+=['','## Competencias por sector','', 'La asociación se obtiene cruzando las celdas respondidas con J, no por cercanía entre columnas. Cada plantilla propuesta contiene las cinco genéricas y cinco específicas. Son borradores sin pesos ni obligatoriedad inventados.','']
for s in sectors:
    lines += [f'### {s["nombre"]} ({s["respuestas"]} respuestas)','']
    for c in competencies:
        if (s['id'],c['id']) in links:
            cols=', '.join(o['columna'] for o in origins if o['competencia_id']==c['id'])
            lines.append(f'- {c["nombre"]} ({cols}).')
    lines.append('')
lines+=['## Competencias pendientes de sector','', 'No hay respuestas para confirmar Administración ni RR. HH. Se sugieren por el contenido, sin crear sectores oficiales ni plantillas para ellos.','']
for cid in unassigned:
    c=next(c for c in competencies if c['id']==cid); oo=[o for o in origins if o['competencia_id']==cid]
    lines.append(f'- {c["nombre"]} ({", ".join(o["columna"] for o in oo)}). {oo[0]["area_sugerida"]}.')
lines+=['','## Limpieza y hallazgos','', '- Se quitaron corchetes, espacios redundantes y puntos finales de nombres. Se separaron nombre y descripción; se conservaron los encabezados originales y las columnas de procedencia.', '- Se normalizaron mayúsculas y acentos de sectores. No hay variantes distintas de sector para fusionar en J.', '- Duplicados textuales unificados:']
for d in duplicates: lines.append(f'  - {", ".join(d["columnas"])} → {d["competencia_id"]}.')
lines+=['- Excel intermedio y avanzado siguen siendo competencias distintas. Tampoco se fusionaron competencias parecidas de bioseguridad, comunicación o ética con distinto alcance.', '- Los 16 encabezados de competencias sin respuestas se conservan como 11 competencias únicas pendientes; no son columnas basura.', '- Se excluyeron 11 marcadores vacíos: BO, BU, BW, BX, BY, BZ, CA, CB, CC, CD y CE.', '- BI, «Columna 60», contiene 72 entradas. No tiene una definición fiable de competencia: queda documentada fuera del catálogo; sus valores originales siguen en el archivo fuente.', '- Faltan dos respuestas específicas: '+', '.join(f'{x["columna"]}{x["fila"]}' for x in missing)+'. No se convierten en cero.', '- No se eliminaron respuestas repetidas de personas: este trabajo depura el catálogo, no decide qué evaluación individual conservar.', '- La escala listada en el JSON contiene solo valores observados. Los niveles no observados no se deducen de este archivo.', '', '## Estructura recomendada para el sistema','', '| Entidad | Campos y relación |', '|---|---|', '| competencias | id, código único, tipo (genérica/específica), nombre, descripción, estado |', '| competencias_genericas / competencias_especificas | Vistas filtradas del catálogo único, evitando duplicar estructura |', '| sectores | id, código único, nombre; los nueve valores de J se conservan como agrupaciones de origen |', '| puestos | id, sector_id, código, nombre; vacío hasta contar con el maestro de puestos |', '| sector_competencias | sector_id + competencia_id, evidencia; clave compuesta única |', '| plantillas_evaluacion | id, código, versión, año, sector_id, puesto_id opcional, estado; única por código y versión |', '| plantilla_competencias | plantilla_id + competencia_id, orden, peso opcional, obligatorio opcional; clave compuesta única |', '| origen_columnas | competencia_id, archivo/huella, hoja, columna, encabezado original |', '', 'Todas las referencias deben tener claves foráneas. Las versiones publicadas deben conservar sus textos y reglas para no alterar evaluaciones históricas. El JSON usa identificadores locales para relacionar registros; no deben confundirse con IDs existentes de la base.', '', 'La fuente mezcla áreas y roles: «Enfermería - Auxiliar de transporte», «Recepción / Cajero» y «Auxiliar de servicios». Se conservan como agrupaciones originales. Separarlos en sectores y puestos requiere contrastar el maestro de personal; no se inventaron puestos.', '', '## Encaje con el proyecto actual','', 'El proyecto ya tiene edd_instrumentos, edd_bloques, edd_items y edd_periodo_instrumentos en la migración 2026_09_25_170000_create_edd_configuration_tables.php. Conviene reutilizar edd_instrumentos como plantillas versionadas, edd_bloques para genéricas/específicas y edd_items para los textos de cada versión, vinculando competencia_codigo con el catálogo. edd_periodo_instrumentos permite aplicar las plantillas por período; sus criterios deberán usar sectores/puestos del maestro validado.', '', 'No ejecutar una segunda estructura de plantillas en paralelo. El catálogo JSON es una propuesta de datos, no una migración aplicada. La integración debe resolver los IDs de sectores/puestos reales, pesos, obligatoriedad y escala completa antes de publicar. No se modificó la aplicación, la base de datos ni el Excel original.', '', '## Verificación','', 'Se verificó cobertura de las 83 columnas, unicidad de competencias y vínculos, referencias del catálogo y nueve plantillas de diez ítems (90 relaciones). El catálogo no incluye nombres, documentos ni calificaciones individuales.']
(OUT/'informe.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
print(json.dumps(dict(genericas=5,especificas=len(payload['competencias_especificas']),sectores=len(sectors),plantillas=len(templates),vinculos=len(items),duplicados=duplicates,faltantes=missing,escala=observed_scale),ensure_ascii=True))
